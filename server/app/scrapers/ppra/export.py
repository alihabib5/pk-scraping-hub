"""Build styled Excel workbook from ppra_bids with proper date/time types and Run Summary."""

import logging
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl.styles import Font
from sqlalchemy import select

from app.core import excel_style
from app.config import settings
from app.db import SessionLocal
from app.scrapers.ppra.filters import FIELD_NAMES, build_search_url, filter_labels
from app.scrapers.ppra.models import PpraBid, PpraRun
from app.scrapers.ppra.schemas import PpraFilters, PpraOptions

logger = logging.getLogger(__name__)

PKT = ZoneInfo("Asia/Karachi")

def _to_pkt_naive(dt: datetime | None) -> datetime | None:
    """Convert tz-aware or UTC datetime to naive Pakistan Standard Time (PKT)."""
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(PKT).replace(tzinfo=None)


def _format_filters_summary(filters: dict[str, Any]) -> str:
    """The applied filters as `Category: Goods · Sector: Info and Comm Tech · …`."""
    parts = [f"{FIELD_NAMES[field]}: {label}" for field, label in filter_labels(filters)]
    adv_from, adv_to = filters.get("advertise_from"), filters.get("advertise_to")
    if adv_from or adv_to:
        parts.append(f"Advertised: {adv_from or 'any'} to {adv_to or 'any'}")
    if filters.get("closing_date"):
        parts.append(f"Closing Date: {filters['closing_date']}")
    return " · ".join(parts) if parts else "All active tenders"


def generate_excel(run_id: str, out_path: str | Path) -> int:
    """Build publication-grade PPRA Tenders Excel workbook from ppra_bids for run_id.

    - Sheet 1 ('PPRA Tenders'):
      Real Excel dates (Advertised, Closing Date), real Excel time (Closing Time PKT),
      live Days Left formula (=MAX(0, ClosingDate-TODAY())), clickable URLs, and Scraped At.
    - Sheet 2 ('Run Summary'):
      Run ID, start/finish times (PKT), applied filters, row counters, and time zone details.
    """
    session = SessionLocal()
    try:
        run: PpraRun | None = session.get(PpraRun, run_id)
        rows = (
            session.execute(
                select(PpraBid)
                .where(PpraBid.run_id == run_id)
                .order_by(PpraBid.page_no, PpraBid.position)
            )
            .scalars()
            .all()
        )
    finally:
        session.close()

    workbook, sheet = excel_style.new_workbook("PPRA Tenders")

    # Scraped At datetime in PKT
    run_start_pkt = _to_pkt_naive(run.started_at if run else None) or datetime.now(PKT).replace(
        tzinfo=None
    )
    run_finish_pkt = _to_pkt_naive(run.finished_at if run else None)

    # Set workbook properties
    workbook.properties.title = "PPRA Tenders"
    workbook.properties.creator = "PK-Hub"
    workbook.properties.created = run_start_pkt

    headers = [
        "Tender No",
        "Title",
        "Ref No",
        "Sector",
        "Organization",
        "Department / Office",
        "City",
        "Status",
        "Advertised Date",
        "Closing Date",
        "Closing Time (PKT)",
        "Days Left",
        "Link",
        "Change",
        "Scraped At (PKT)",
    ]

    sheet.append([excel_style.sanitize_cell(h) for h in headers])

    adv_count = 0
    closing_count = 0

    link_font = Font(color="0563C1", underline="single")

    for idx, bid in enumerate(rows, start=2):  # Row 1 is header, data starts at 2
        # 1. Advertised Date (real date)
        adv_val: date | None = None
        if bid.advertised_on:
            adv_val = bid.advertised_on
            adv_count += 1

        # 2. Closing Date and Time (split from closing_at)
        closing_date_val: date | None = None
        closing_time_val: time | None = None
        if bid.closing_at:
            closing_pkt = _to_pkt_naive(bid.closing_at)
            if closing_pkt:
                closing_date_val = closing_pkt.date()
                closing_time_val = closing_pkt.time().replace(microsecond=0)
                closing_count += 1

        # 3. Days Left (Excel formula referencing Column J = Closing Date)
        # In Excel: J is column 10
        if closing_date_val is not None:
            days_left_formula = f"=MAX(0, J{idx}-TODAY())"
        else:
            days_left_formula = ""

        # Row values
        row_values = [
            bid.tender_no or "",
            bid.title or "",
            bid.tender_ref_no or "",
            bid.sector or "",
            bid.organization or "",
            bid.department or "",
            bid.city or "",
            bid.status or "Published",
            adv_val,
            closing_date_val,
            closing_time_val,
            days_left_formula if days_left_formula else None,
            bid.detail_url or "",
            bid.change_status or "new",
            run_start_pkt,
        ]

        sheet.append(row_values)

        # Format cell types and styles for the appended row
        # Col I (9): Advertised Date
        adv_cell = sheet.cell(row=idx, column=9)
        if adv_val is not None:
            adv_cell.number_format = "yyyy-mm-dd"

        # Col J (10): Closing Date
        close_d_cell = sheet.cell(row=idx, column=10)
        if closing_date_val is not None:
            close_d_cell.number_format = "yyyy-mm-dd"

        # Col K (11): Closing Time
        close_t_cell = sheet.cell(row=idx, column=11)
        if closing_time_val is not None:
            close_t_cell.number_format = "hh:mm"

        # Col L (12): Days Left
        days_cell = sheet.cell(row=idx, column=12)
        if days_left_formula:
            days_cell.number_format = "0"

        # Col M (13): Link (hyperlink)
        link_cell = sheet.cell(row=idx, column=13)
        if bid.detail_url:
            link_cell.hyperlink = bid.detail_url
            link_cell.font = link_font

        # Col O (15): Scraped At
        scraped_cell = sheet.cell(row=idx, column=15)
        scraped_cell.number_format = "yyyy-mm-dd hh:mm"

    # Style header row, freeze panes, auto-fit columns and enable auto-filter
    excel_style.format_excel_headers(sheet, header_row=1, autofit=True, freeze=True, auto_filter=True)

    # -------------------------------------------------------------------------
    # Sheet 2: Run Summary
    # -------------------------------------------------------------------------
    summary_sheet = workbook.create_sheet(title="Run Summary")
    summary_headers = ["Parameter", "Value"]
    summary_sheet.append([excel_style.sanitize_cell(h) for h in summary_headers])

    filters_data = run.filters if run and run.filters else {}
    portal_url = settings.ppra_base_url.rstrip("/") + settings.ppra_active_tenders_path
    search_url = build_search_url(PpraFilters(**filters_data), portal_url)
    run_options = PpraOptions(**(run.options if run and run.options else {}))
    runway = run_options.get_runway_days()
    skipped = run.skipped_closing if run else 0
    unverified = run.kept_unverified if run else 0
    closing_guard = (
        "Off — all tenders kept"
        if runway is None
        else f"{runway} days (skipped closing sooner: {skipped}, kept with unreadable date: {unverified})"
    )
    filters_desc = _format_filters_summary(filters_data)

    duration_str = "—"
    if run_start_pkt and run_finish_pkt:
        sec = int((run_finish_pkt - run_start_pkt).total_seconds())
        m, s = divmod(sec, 60)
        duration_str = f"{m}m {s}s" if m > 0 else f"{s}s"

    summary_rows = [
        ("Run ID", run_id),
        ("Status", (run.status if run else "completed").capitalize()),
        ("Scrape Started (PKT)", run_start_pkt.strftime("%Y-%m-%d %H:%M:%S") if run_start_pkt else "—"),
        ("Scrape Finished (PKT)", run_finish_pkt.strftime("%Y-%m-%d %H:%M:%S") if run_finish_pkt else "—"),
        ("Duration", duration_str),
        ("Source Portal", portal_url),
        ("Search Filters Applied", filters_desc),
        ("Search URL", search_url),
        ("Closing Guard", closing_guard),
        ("Karachi Tenders", "Excluded" if run_options.exclude_karachi else "Included"),
        ("Total Tenders Scraped", len(rows)),
        ("Total Expected on Portal", run.total_expected if (run and run.total_expected) else len(rows)),
        ("New Tenders", run.rows_new if run else len(rows)),
        ("Changed Tenders", run.rows_changed if run else 0),
        ("Time Zone Note", "All timestamps and tender closing times are in Pakistan Standard Time (PKT, UTC+5)"),
    ]

    for param, val in summary_rows:
        summary_sheet.append([param, str(val) if val is not None else "—"])

    excel_style.format_excel_headers(summary_sheet, header_row=1, autofit=True, freeze=True, auto_filter=False)

    out_file = Path(out_path)
    out_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(out_file))

    logger.info(
        "[run %s] generated PPRA Excel with %d rows at %s (dates parsed: advertised %d/%d, closing %d/%d, missing %d)",
        run_id,
        len(rows),
        out_file,
        adv_count,
        len(rows),
        closing_count,
        len(rows),
        len(rows) - closing_count,
    )
    return len(rows)
