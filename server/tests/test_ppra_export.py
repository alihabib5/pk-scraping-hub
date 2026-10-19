"""Tests for PPRA Excel export: real dates, PKT timezone, formulas, summary sheet, and filename."""

import re
from datetime import date, datetime, time, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
import pytest

from app.core import exports
from app.db import Base, SessionLocal, engine
from app.scrapers.ppra import export
from app.scrapers.ppra.models import PpraBid, PpraRun

PKT = ZoneInfo("Asia/Karachi")


@pytest.fixture(autouse=True)
def setup_db():
    Base.metadata.create_all(bind=engine)
    yield


def test_ppra_filename_format():
    """Verify filename matches ^PPRA_Tenders_\\d{4}-\\d{2}-\\d{2}_\\d{4}\\.xlsx$."""
    run = {
        "scraper": "ppra",
        "run_id": "test_run_123",
        "started_at": "2026-10-14T19:30:00Z",  # 19:30 UTC = 00:30 next day in PKT
    }
    filename = exports.excel_name(run)
    assert re.match(r"^PPRA_Tenders_\d{4}-\d{2}-\d{2}_\d{4}\.xlsx$", filename), f"Got: {filename}"
    assert filename == "PPRA_Tenders_2026-10-15_0030.xlsx"


@pytest.mark.parametrize(
    "filters, expected",
    [
        (
            {"advertise_from": "2026-10-13", "advertise_to": "2026-10-15", "category": "Goods",
             "country": "Pakistan", "sector": "ALL"},
            "PPRA_Tenders(13-10-2026 to 15-10-2026, Goods).xlsx",
        ),
        (
            {"advertise_from": "2026-10-16", "advertise_to": "2026-10-17", "category": "Goods",
             "tender_type": "Tender Notice", "city": "Lahore"},
            "PPRA_Tenders(16-10-2026 to 17-10-2026, Goods, Tender Notice, Lahore).xlsx",
        ),
        ({"category": "Works"}, "PPRA_Tenders(Works).xlsx"),
        ({"advertise_from": "2026-10-13", "advertise_to": "2026-10-13"}, "PPRA_Tenders(13-10-2026).xlsx"),
        ({"advertise_from": "2026-10-13"}, "PPRA_Tenders(from 13-10-2026).xlsx"),
        ({"advertise_to": "2026-10-15"}, "PPRA_Tenders(to 15-10-2026).xlsx"),
    ],
)
def test_ppra_filename_uses_advertise_dates(filters, expected):
    run = {"scraper": "ppra", "run_id": "r1", "started_at": "2026-10-17T06:50:03", "filters": filters}
    assert exports.excel_name(run) == expected


def test_ppra_filename_naive_start_is_local_time():
    """run_manager stores started_at as naive local time; it must not be shifted as if UTC."""
    started = datetime(2026, 10, 17, 6, 50).astimezone()  # local wall clock, made aware
    run = {"scraper": "ppra", "run_id": "r1", "started_at": "2026-10-17T06:50:03", "filters": {}}
    expected = started.astimezone(PKT).strftime("PPRA_Tenders_%Y-%m-%d_%H%M.xlsx")
    assert exports.excel_name(run) == expected


def test_ppra_export_full(tmp_path: Path):
    """Test full workbook generation: date types, formulas, timezone shift, and summary sheet."""
    session = SessionLocal()
    run_id = f"test_run_{datetime.now().timestamp()}"

    start_utc = datetime(2026, 10, 14, 14, 30, 0, tzinfo=timezone.utc)  # 19:30 PKT
    run = PpraRun(
        run_id=run_id,
        status="completed",
        filters={"category": "1", "keyword": "Laptop", "advertise_from": "2026-10-14"},
        filters_hash="abc123hash",
        options={},
        started_at=start_utc,
        finished_at=datetime(2026, 10, 14, 14, 35, 0, tzinfo=timezone.utc),
        total_expected=2,
        rows_scraped=2,
        rows_new=2,
        rows_changed=0,
        skipped_closing=1,
        kept_unverified=0,
    )
    session.add(run)

    # Tender 1: valid dates
    bid1 = PpraBid(
        run_id=run_id,
        page_no=1,
        position=1,
        tender_no="TS1001",
        title="Procurement of IT Equipment",
        tender_ref_no="IT-2026-01",
        sector="Information Technology",
        organization="Federal Board of Revenue",
        department="IT Division",
        city="Islamabad",
        country="Pakistan",
        status="Published",
        advertised_on=date(2026, 10, 14),
        closing_at=datetime(2026, 11, 19, 14, 30, 0, tzinfo=timezone.utc),  # 19:30 PKT
        detail_url="https://epms.ppra.gov.pk/public/tenders/tender-details/1001",
        content_hash="hash1",
        change_status="new",
    )

    # Tender 2: missing dates (tests robustness)
    bid2 = PpraBid(
        run_id=run_id,
        page_no=1,
        position=2,
        tender_no="TS1002",
        title="Civil Works",
        tender_ref_no="CW-02",
        sector="Construction",
        organization="NHA",
        department=None,
        city="Lahore",
        country="Pakistan",
        status="Published",
        advertised_on=None,
        closing_at=None,
        detail_url=None,
        content_hash="hash2",
        change_status="new",
    )

    session.add_all([bid1, bid2])
    session.commit()
    session.close()

    out_file = tmp_path / "test_export.xlsx"
    count = export.generate_excel(run_id, out_file)
    assert count == 2
    assert out_file.is_file()

    # Read back workbook using openpyxl
    wb = openpyxl.load_workbook(str(out_file), data_only=False)

    # 1. Sheets check
    assert "PPRA Tenders" in wb.sheetnames
    assert "Run Summary" in wb.sheetnames

    # 2. Sheet 1: PPRA Tenders
    ws = wb["PPRA Tenders"]

    # Header check
    headers = [ws.cell(row=1, column=col).value for col in range(1, 16)]
    assert "Advertised Date" in headers
    assert "Closing Date" in headers
    assert "Closing Time (PKT)" in headers
    assert "Days Left" in headers
    assert "Scraped At (PKT)" in headers

    # Auto-filter check
    assert ws.auto_filter.ref is not None
    assert "A1:O3" in ws.auto_filter.ref

    # Row 2 (bid1) assertions
    adv_cell = ws.cell(row=2, column=9)
    assert isinstance(adv_cell.value, (date, datetime))
    assert (adv_cell.value.date() if isinstance(adv_cell.value, datetime) else adv_cell.value) == date(2026, 10, 14)
    assert adv_cell.number_format == "yyyy-mm-dd"

    # Closing Date: 2026-11-19
    close_d_cell = ws.cell(row=2, column=10)
    assert isinstance(close_d_cell.value, (date, datetime))
    assert (close_d_cell.value.date() if isinstance(close_d_cell.value, datetime) else close_d_cell.value) == date(2026, 11, 19)
    assert close_d_cell.number_format == "yyyy-mm-dd"

    # Closing Time: 19:30 (14:30 UTC + 5h = 19:30 PKT)
    close_t_cell = ws.cell(row=2, column=11)
    assert close_t_cell.value == time(19, 30, 0)
    assert close_t_cell.number_format == "hh:mm"

    # Days Left: Formula
    days_cell = ws.cell(row=2, column=12)
    assert days_cell.value == "=MAX(0, J2-TODAY())"
    assert days_cell.number_format == "0"

    # Hyperlink check
    link_cell = ws.cell(row=2, column=13)
    assert link_cell.hyperlink is not None
    assert link_cell.hyperlink.target == "https://epms.ppra.gov.pk/public/tenders/tender-details/1001"

    # Scraped At check
    scraped_cell = ws.cell(row=2, column=15)
    assert isinstance(scraped_cell.value, datetime)
    assert scraped_cell.value.year == 2026
    assert scraped_cell.value.hour == 19  # 14:30 UTC = 19:30 PKT
    assert scraped_cell.value.minute == 30
    assert scraped_cell.number_format == "yyyy-mm-dd hh:mm"

    # Row 3 (bid2 - missing dates) assertions
    adv2_cell = ws.cell(row=3, column=9)
    assert adv2_cell.value is None

    close2_d_cell = ws.cell(row=3, column=10)
    assert close2_d_cell.value is None

    close2_t_cell = ws.cell(row=3, column=11)
    assert close2_t_cell.value is None

    days2_cell = ws.cell(row=3, column=12)
    assert days2_cell.value is None

    # 3. Sheet 2: Run Summary
    ws_sum = wb["Run Summary"]
    summary_data = {
        ws_sum.cell(row=r, column=1).value: ws_sum.cell(row=r, column=2).value
        for r in range(2, ws_sum.max_row + 1)
    }

    assert summary_data["Run ID"] == run_id
    assert summary_data["Status"] == "Completed"
    assert "Goods" in summary_data["Search Filters Applied"]  # category 1 mapped to Goods
    assert "Laptop" in summary_data["Search Filters Applied"]
    assert "PKT, UTC+5" in summary_data["Time Zone Note"]
    assert summary_data["Total Tenders Scraped"] == "2"

    wb.close()
